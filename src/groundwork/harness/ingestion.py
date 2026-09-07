"""Convert raw source documents into anchored markdown for downstream LLM steps.

Format conversion is industry-agnostic, so it lives in the harness. Callers are
responsible for whatever domain meaning they attach to the extracted text.

The single entry point is :func:`extract_document`. Every extracted document
carries *anchors* — inline HTML comments such as ``<!-- page: 3 -->`` — so a
downstream extraction step can map any span of text back to its location in the
source file. Anchors are HTML comments rather than headings so they never
collide with real document structure and are trivially strippable via
:data:`ANCHOR_PATTERN`.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Literal

from groundwork.harness.cleanup import (
    clean_markdown,
    flatten_table_cell,
    rows_to_pipe_table,
)

if TYPE_CHECKING:  # pragma: no cover - import-time typing only
    from docx.table import Table
    from docx.text.paragraph import Paragraph

__all__ = [
    "ANCHOR_PATTERN",
    "DEFAULT_PARAGRAPHS_PER_ANCHOR",
    "SUPPORTED_SUFFIXES",
    "DocumentReadError",
    "ExtractedDocument",
    "IngestionError",
    "OutputPathError",
    "UnsupportedFormatError",
    "extract_document",
]

PDF_SUFFIX = ".pdf"
DOCX_SUFFIX = ".docx"
SUPPORTED_SUFFIXES: frozenset[str] = frozenset({PDF_SUFFIX, DOCX_SUFFIX})

PAGE_ANCHOR = "<!-- page: {index} -->"
PARAGRAPH_ANCHOR = "<!-- paragraph: {index} -->"

#: Matches any anchor emitted by this module; group ``kind`` is ``page`` or
#: ``paragraph``, group ``index`` is the 1-based page number or the 0-based
#: index into ``docx.Document.paragraphs``.
ANCHOR_PATTERN = re.compile(r"<!-- (?P<kind>page|paragraph): (?P<index>\d+) -->")

#: How many rendered DOCX blocks may pass before a positional anchor is forced.
#: Headings always emit an anchor; this bounds anchor spacing in documents that
#: use few or no heading styles.
DEFAULT_PARAGRAPHS_PER_ANCHOR = 25

_MAX_HEADING_LEVEL = 6


class IngestionError(Exception):
    """Base class for every error raised by this module."""


class UnsupportedFormatError(IngestionError):
    """The file extension is not one this module knows how to extract."""


class DocumentReadError(IngestionError):
    """The file exists but could not be parsed (corrupt, encrypted, truncated)."""


class OutputPathError(IngestionError):
    """The requested ``output_path`` is unusable (collides with input, unwritable)."""


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    """Result of an extraction run.

    Attributes:
        source_path: The file the content was read from.
        source_format: ``"pdf"`` or ``"docx"``.
        anchor_kind: ``"page"`` for PDFs, ``"paragraph"`` for DOCX.
        anchor_count: Number of anchors embedded in ``markdown``.
        markdown: The extracted content, with anchors inline.
        output_path: Where the markdown was cached, if ``output_path`` was given.
    """

    source_path: Path
    source_format: Literal["pdf", "docx"]
    anchor_kind: Literal["page", "paragraph"]
    anchor_count: int
    markdown: str
    output_path: Path | None = None

    def __str__(self) -> str:
        return self.markdown


def extract_document(
    input_path: str | Path,
    *,
    output_path: str | Path | None = None,
    paragraphs_per_anchor: int = DEFAULT_PARAGRAPHS_PER_ANCHOR,
    clean: bool = True,
) -> ExtractedDocument:
    """Extract a PDF or DOCX file as anchored markdown.

    PDFs are read with ``pymupdf4llm``, which reconstructs headings, lists and
    tables as markdown instead of emitting the flat character dump that
    ``pymupdf.Page.get_text()`` produces. DOCX files are read with
    ``python-docx``, because PyMuPDF's Office support is lossy; body paragraphs
    and tables are walked in document order so table content is not dropped.

    Traceability differs by format, and the difference is inherent to the
    formats rather than a shortcut:

    * **PDF** — pages are a real structural unit, so each page's content is
      prefixed with ``<!-- page: N -->`` (1-based). Empty pages still emit their
      anchor, so numbering stays continuous.
    * **DOCX** — a ``.docx`` stores a reflowable stream of blocks; page breaks
      are computed by the renderer from font metrics, paper size and printer
      settings, so there is no page number to report without laying the document
      out. Anchors are therefore ``<!-- paragraph: N -->``, where ``N`` indexes
      into ``docx.Document.paragraphs``. An anchor is emitted before every
      heading and at least every ``paragraphs_per_anchor`` rendered blocks, so a
      span maps back to the nearest preceding anchor rather than to an exact
      page. Callers needing true DOCX page numbers must convert to PDF first.

    The input file is only ever opened for reading and is never modified.

    Args:
        input_path: Path to a ``.pdf`` or ``.docx`` file (case-insensitive).
        output_path: Optional path to also cache the markdown to, UTF-8 encoded.
            Parent directories are created. Must not be ``input_path``.
        paragraphs_per_anchor: Maximum rendered DOCX blocks between positional
            anchors. Ignored for PDFs. Must be at least 1.
        clean: Whether to pass the extracted text through
            :func:`groundwork.harness.cleanup.clean_markdown`, which removes
            repeating page furniture, normalizes symbol-font bullets and
            converts the HTML tables this function requests into pipe tables.
            Set to ``False`` to inspect raw extractor output.

    Returns:
        An :class:`ExtractedDocument`; use its ``markdown`` attribute (or
        ``str()`` it) for the text.

    Raises:
        FileNotFoundError: ``input_path`` does not exist or is not a file.
        UnsupportedFormatError: The extension is neither ``.pdf`` nor ``.docx``.
        DocumentReadError: The file could not be parsed.
        OutputPathError: ``output_path`` equals ``input_path``, or writing failed.
        ValueError: ``paragraphs_per_anchor`` is less than 1.
    """
    if paragraphs_per_anchor < 1:
        raise ValueError(
            f"paragraphs_per_anchor must be at least 1, got {paragraphs_per_anchor}"
        )

    source = Path(input_path).expanduser()
    if not source.is_file():
        raise FileNotFoundError(f"No such document to ingest: {source}")

    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise UnsupportedFormatError(
            f"Cannot ingest {source.name!r}: extension {suffix or '(none)'!r} is not "
            f"supported. Supported extensions: {supported}."
        )

    destination = _resolve_output_path(source, output_path)

    if suffix == PDF_SUFFIX:
        markdown, anchor_count = _extract_pdf(source)
        source_format: Literal["pdf", "docx"] = "pdf"
        anchor_kind: Literal["page", "paragraph"] = "page"
    else:
        markdown, anchor_count = _extract_docx(source, paragraphs_per_anchor)
        source_format = "docx"
        anchor_kind = "paragraph"

    if clean:
        markdown = clean_markdown(markdown)

    if destination is not None:
        _write_markdown(destination, markdown)

    return ExtractedDocument(
        source_path=source,
        source_format=source_format,
        anchor_kind=anchor_kind,
        anchor_count=anchor_count,
        markdown=markdown,
        output_path=destination,
    )


def _resolve_output_path(source: Path, output_path: str | Path | None) -> Path | None:
    """Validate ``output_path`` against ``source``, refusing to overwrite the input."""
    if output_path is None:
        return None

    destination = Path(output_path).expanduser()
    # Compare resolved paths so symlinks and `./`-style aliases are caught too.
    if destination.resolve() == source.resolve():
        raise OutputPathError(
            f"output_path must differ from input_path; both resolve to {source.resolve()}. "
            "Writing there would truncate the source document."
        )
    return destination


def _write_markdown(destination: Path, markdown: str) -> None:
    """Cache ``markdown`` to ``destination`` as UTF-8, creating parent directories."""
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        with destination.open("w", encoding="utf-8") as handle:
            handle.write(markdown)
    except OSError as exc:
        raise OutputPathError(f"Could not write extracted markdown to {destination}: {exc}") from exc


def _extract_pdf(source: Path) -> tuple[str, int]:
    """Return ``(markdown, anchor_count)`` for a PDF, one anchor per page."""
    # Imported lazily: the PDF stack is heavy, and importing it at module scope
    # would make `import groundwork.harness.ingestion` slow for DOCX-only callers.
    import pymupdf
    import pymupdf4llm

    try:
        with pymupdf.open(source) as document:
            if document.needs_pass:
                raise DocumentReadError(
                    f"{source.name} is password-protected; decrypt it before ingestion."
                )
            # table_output="html" is not cosmetic. pymupdf4llm's layout engine
            # writes markdown table cells by concatenating spans without their
            # inter-word gaps, so "(if applicable)" arrives as "(ifapplicable)";
            # its HTML writer preserves them. Prose is unaffected either way.
            # `clean_markdown` converts the HTML back to pipe tables downstream.
            chunks = pymupdf4llm.to_markdown(
                document, page_chunks=True, table_output="html"
            )
    except DocumentReadError:
        raise
    except Exception as exc:  # pymupdf raises bare RuntimeError/FileDataError subtypes
        raise DocumentReadError(
            f"Could not extract text from PDF {source}: {exc}"
        ) from exc

    sections: list[str] = []
    for position, chunk in enumerate(chunks, start=1):
        # Layout mode reports "page_number", legacy mode reports "page".
        metadata = chunk.get("metadata", {})
        page_number = metadata.get("page_number", metadata.get("page", position))
        anchor = PAGE_ANCHOR.format(index=page_number)
        text = (chunk.get("text") or "").strip()
        sections.append(f"{anchor}\n\n{text}" if text else anchor)

    return "\n\n".join(sections), len(sections)


def _extract_docx(source: Path, paragraphs_per_anchor: int) -> tuple[str, int]:
    """Return ``(markdown, anchor_count)`` for a DOCX, anchored by paragraph index."""
    import docx
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    sections: list[str] = []
    anchor_count = 0
    paragraph_index = 0
    blocks_since_anchor = paragraphs_per_anchor  # force an anchor before the first block

    try:
        with source.open("rb") as handle:
            document = docx.Document(handle)
            # Walk the body element directly: `document.paragraphs` skips tables,
            # and tables carry most of the structured content in real-world packs.
            for element in document.element.body.iterchildren():
                if isinstance(element, CT_P):
                    current_index = paragraph_index
                    paragraph_index += 1
                    paragraph = Paragraph(element, document)
                    rendered = _render_paragraph(paragraph)
                    is_heading = rendered.startswith("#")
                elif isinstance(element, CT_Tbl):
                    current_index = paragraph_index
                    rendered = _render_table(Table(element, document))
                    is_heading = False
                else:
                    continue

                if not rendered:
                    continue

                if is_heading or blocks_since_anchor >= paragraphs_per_anchor:
                    sections.append(PARAGRAPH_ANCHOR.format(index=current_index))
                    anchor_count += 1
                    blocks_since_anchor = 0

                sections.append(rendered)
                blocks_since_anchor += 1
    except IngestionError:
        raise
    except Exception as exc:  # python-docx raises PackageNotFoundError, KeyError, ...
        raise DocumentReadError(
            f"Could not extract text from DOCX {source}: {exc}"
        ) from exc

    return "\n\n".join(sections), anchor_count


def _render_paragraph(paragraph: Paragraph) -> str:
    """Render one DOCX paragraph as markdown, or ``""`` if it holds no text."""
    text = paragraph.text.strip()
    if not text:
        return ""

    style_name = paragraph.style.name if paragraph.style is not None else ""
    style_name = style_name or ""

    if style_name == "Title":
        return f"# {text}"
    if style_name.startswith("Heading"):
        return f"{'#' * _heading_level(style_name)} {text}"
    if style_name.startswith("List Bullet"):
        return f"- {text}"
    if style_name.startswith("List Number"):
        return f"1. {text}"
    return text


def _heading_level(style_name: str) -> int:
    """Map a ``Heading N`` style name to a markdown heading depth, clamped to 1..6."""
    digits = re.search(r"(\d+)", style_name)
    if digits is None:
        return 1
    return min(max(int(digits.group(1)), 1), _MAX_HEADING_LEVEL)


def _render_table(table: Table) -> str:
    """Render a DOCX table as a markdown pipe table, or ``""`` if it is empty."""
    return rows_to_pipe_table(
        [[flatten_table_cell(cell.text) for cell in row.cells] for row in table.rows]
    )
