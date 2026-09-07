"""Tests for :mod:`groundwork.harness.ingestion`.

Fixtures build their own tiny PDF/DOCX files so the suite does not depend on the
sample tender pack under ``vertical/data/`` — the harness must stay testable
without any vertical's data.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from groundwork.harness.ingestion import (
    ANCHOR_PATTERN,
    DocumentReadError,
    ExtractedDocument,
    IngestionError,
    OutputPathError,
    UnsupportedFormatError,
    extract_document,
)

PDF_PAGE_COUNT = 3


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """A three-page PDF with distinct, searchable text on each page."""
    import pymupdf

    path = tmp_path / "sample.pdf"
    with pymupdf.open() as document:
        for number in range(1, PDF_PAGE_COUNT + 1):
            page = document.new_page()
            page.insert_text((72, 100), f"Requirement {number}: supply widgets.", fontsize=14)
        document.save(path)
    return path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """A DOCX with a heading, body paragraphs and a table."""
    import docx

    path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_heading("Scope of Works", level=1)
    document.add_paragraph("The supplier shall deliver widgets.")
    document.add_heading("Evaluation", level=2)
    document.add_paragraph("Bids are scored on price and quality.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Criterion"
    table.cell(0, 1).text = "Weighting"
    table.cell(1, 0).text = "Price"
    table.cell(1, 1).text = "60%"

    document.save(path)
    return path


def _fingerprint(path: Path) -> tuple[bytes, int]:
    return hashlib.sha256(path.read_bytes()).digest(), path.stat().st_size


def test_pdf_output_embeds_a_marker_for_every_page(sample_pdf: Path) -> None:
    result = extract_document(sample_pdf)

    assert result.anchor_kind == "page"
    assert result.anchor_count == PDF_PAGE_COUNT
    for number in range(1, PDF_PAGE_COUNT + 1):
        assert f"<!-- page: {number} -->" in result.markdown

    anchors = [match.group("index") for match in ANCHOR_PATTERN.finditer(result.markdown)]
    assert anchors == ["1", "2", "3"], "page anchors must appear in document order"


def test_pdf_page_text_follows_its_own_marker(sample_pdf: Path) -> None:
    markdown = extract_document(sample_pdf).markdown

    for number in range(1, PDF_PAGE_COUNT + 1):
        marker_at = markdown.index(f"<!-- page: {number} -->")
        text_at = markdown.index(f"Requirement {number}")
        assert marker_at < text_at
        if number < PDF_PAGE_COUNT:
            assert text_at < markdown.index(f"<!-- page: {number + 1} -->")


@pytest.mark.parametrize("fixture_name", ["sample_pdf", "sample_docx"])
def test_extraction_does_not_mutate_the_input_file(
    fixture_name: str, request: pytest.FixtureRequest
) -> None:
    source: Path = request.getfixturevalue(fixture_name)
    before = _fingerprint(source)

    extract_document(source)

    assert _fingerprint(source) == before


def test_writing_output_does_not_mutate_the_input_file(
    sample_pdf: Path, tmp_path: Path
) -> None:
    before = _fingerprint(sample_pdf)

    result = extract_document(sample_pdf, output_path=tmp_path / "out" / "sample.md")

    assert _fingerprint(sample_pdf) == before
    assert result.output_path is not None
    assert result.output_path.read_text(encoding="utf-8") == result.markdown


def test_unsupported_extension_raises(tmp_path: Path) -> None:
    source = tmp_path / "notes.txt"
    source.write_text("plain text", encoding="utf-8")

    with pytest.raises(UnsupportedFormatError, match=r"\.txt"):
        extract_document(source)


def test_extensionless_file_raises_unsupported_format(tmp_path: Path) -> None:
    source = tmp_path / "tender"
    source.write_bytes(b"%PDF-1.4")

    with pytest.raises(UnsupportedFormatError):
        extract_document(source)


def test_extension_matching_is_case_insensitive(sample_pdf: Path) -> None:
    # Distinct stem, so this is a real second file even on a case-insensitive FS.
    upper = sample_pdf.parent / "SHOUTING.PDF"
    upper.write_bytes(sample_pdf.read_bytes())

    assert extract_document(upper).anchor_count == PDF_PAGE_COUNT


def test_missing_file_raises_file_not_found(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        extract_document(tmp_path / "absent.pdf")


def test_corrupt_pdf_raises_document_read_error(tmp_path: Path) -> None:
    source = tmp_path / "corrupt.pdf"
    source.write_bytes(b"%PDF-1.7\nnot actually a pdf\n")

    with pytest.raises(DocumentReadError):
        extract_document(source)


def test_corrupt_docx_raises_document_read_error(tmp_path: Path) -> None:
    source = tmp_path / "corrupt.docx"
    source.write_bytes(b"PK\x03\x04 truncated zip")

    with pytest.raises(DocumentReadError):
        extract_document(source)


def test_output_path_equal_to_input_path_raises(sample_pdf: Path) -> None:
    before = _fingerprint(sample_pdf)

    with pytest.raises(OutputPathError):
        extract_document(sample_pdf, output_path=sample_pdf)

    assert _fingerprint(sample_pdf) == before, "input must not be truncated"


def test_output_path_symlinked_to_the_input_raises(sample_pdf: Path, tmp_path: Path) -> None:
    """A symlink is a different path string but the same bytes on disk."""
    alias = tmp_path / "alias.pdf"
    alias.symlink_to(sample_pdf)
    before = _fingerprint(sample_pdf)

    with pytest.raises(OutputPathError):
        extract_document(sample_pdf, output_path=alias)

    assert _fingerprint(sample_pdf) == before


def test_docx_uses_paragraph_anchors_and_keeps_structure(sample_docx: Path) -> None:
    result = extract_document(sample_docx)

    assert result.anchor_kind == "paragraph"
    assert result.anchor_count >= 2, "each heading should emit its own anchor"
    assert "<!-- page:" not in result.markdown, "DOCX has no reliable page boundaries"

    assert "# Scope of Works" in result.markdown
    assert "## Evaluation" in result.markdown
    assert "| Criterion | Weighting |" in result.markdown
    assert "| Price | 60% |" in result.markdown

    anchors = [int(m.group("index")) for m in ANCHOR_PATTERN.finditer(result.markdown)]
    assert anchors == sorted(anchors), "paragraph anchors must be monotonic"
    assert all(m.group("kind") == "paragraph" for m in ANCHOR_PATTERN.finditer(result.markdown))


def test_docx_paragraph_anchor_indexes_the_source_paragraph(sample_docx: Path) -> None:
    import docx

    markdown = extract_document(sample_docx).markdown
    match = ANCHOR_PATTERN.search(markdown)
    assert match is not None

    document = docx.Document(str(sample_docx))
    assert document.paragraphs[int(match.group("index"))].text == "Scope of Works"


def test_docx_anchor_stride_bounds_the_gap_between_anchors(tmp_path: Path) -> None:
    import docx

    path = tmp_path / "flat.docx"
    document = docx.Document()
    for number in range(20):
        document.add_paragraph(f"Clause {number}.")
    document.save(path)

    result = extract_document(path, paragraphs_per_anchor=5)

    assert result.anchor_count == 4


def test_invalid_paragraphs_per_anchor_raises(sample_docx: Path) -> None:
    with pytest.raises(ValueError):
        extract_document(sample_docx, paragraphs_per_anchor=0)


def test_result_is_a_string_like_structured_object(sample_pdf: Path) -> None:
    result = extract_document(sample_pdf)

    assert isinstance(result, ExtractedDocument)
    assert str(result) == result.markdown
    assert result.source_format == "pdf"
    assert result.source_path == sample_pdf
    assert result.output_path is None


def test_module_errors_share_a_base_class() -> None:
    for error in (UnsupportedFormatError, DocumentReadError, OutputPathError):
        assert issubclass(error, IngestionError)
